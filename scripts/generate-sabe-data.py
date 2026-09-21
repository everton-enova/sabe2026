"""Generate the application lookup data from the SABE workbook export."""

from __future__ import annotations

import json
import sys
import unicodedata
from pathlib import Path

import openpyxl


def key(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(char for char in text if not unicodedata.combining(char)).upper().split())


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def nte_label(value: object) -> str:
    digits = "".join(char for char in clean(value) if char.isdigit())
    return f"NTE {int(digits):02d}"


def main() -> None:
    source = Path(sys.argv[1] if len(sys.argv) > 1 else "work/sabe-current.xlsx")
    target = Path(sys.argv[2] if len(sys.argv) > 2 else "src/data/sabe.json")
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)

    municipality_nte: dict[str, str] = {}
    for row in workbook["MUN-NTE"].iter_rows(min_row=2, values_only=True):
        values = list(row) + [None, None]
        if values[0] and values[1]:
            municipality_nte[key(values[0])] = nte_label(values[1])
    # Confirmed manually in the source spreadsheet.
    municipality_nte["BARRO PRETO"] = "NTE 05"

    locations = []
    for row in workbook["MUNICIPIOS POR POLOS"].iter_rows(min_row=2, values_only=True):
        values = list(row) + [None, None]
        if not values[0] or not values[1]:
            continue
        municipality = clean(values[1])
        if key(municipality) == "GOVERNADOR LOMANTO JUNIOR":
            municipality = "BARRO PRETO"
        locations.append(
            {
                "nte": municipality_nte[key(municipality)],
                "polo": clean(values[0]).upper(),
                "municipio": municipality.upper(),
            }
        )

    coordinators = []
    for row in workbook["CP- SABE "].iter_rows(min_row=2, values_only=True):
        values = list(row) + [None] * 16
        if not values[2]:
            continue
        location = {"nte": nte_label(values[1]), "polo": clean(values[2]).upper()}
        coordinators.append(location)

    payload = {"locations": locations, "coordinators": coordinators}
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Generated {len(locations)} location rows and {len(coordinators)} coordinator locations in {target}")


if __name__ == "__main__":
    main()
